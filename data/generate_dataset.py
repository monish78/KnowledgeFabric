"""Generate the Graphbase test dataset.

Deterministic (fixed seed): running it twice gives identical files.  Writes
  data/samples/   the files a user would upload
  data/expected/manifest.json   ground truth used by the automated tests

The data is deliberately messy.  The ingest contract the backend must honour
(also written into the manifest) is:
  * header rows may be offset by title/blank rows; blank rows are ignored
  * headers are matched case/space/underscore-insensitively
  * key values are compared after strip() + upper()
  * numbers may carry currency symbols, thousands separators or accounting
    parentheses; dates may be ISO, DD/MM/YYYY, "Sep 22, 2026" or Excel dates;
    booleans may be Y/N/yes/no/TRUE/FALSE/1/0
  * all nodes are created before relationships (forward references are legal)
  * a row is rejected (and reported) if a node key it needs is missing or a
    relationship endpoint it references does not exist; rejection is per row
  * duplicate rows merge on the key property; for repeated relationship rows
    the last row wins
"""
import csv
import datetime as dt
import json
import random
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).parent
SAMPLES = ROOT / "samples"
EXPECTED = ROOT / "expected"
rng = random.Random(20260929)

FIRST = ["Aarav", "Ananya", "Vikram", "Meera", "Rohan", "Sneha", "Arjun", "Kavya", "Karthik", "Divya",
         "Rahul", "Priyanka", "Suresh", "Lakshmi", "Nikhil", "Pooja", "Ravi", "Deepa", "Sanjay", "Nandini",
         "Farhan", "Zoya", "Gurpreet", "Harleen", "Joseph", "Mary", "Tenzin", "Anjali", "Venkat", "Shreya"]
LAST = ["Sharma", "Iyer", "Nair", "Reddy", "Menon", "Gupta", "Krishnan", "Pillai", "Das", "Mehta",
        "Kulkarni", "Banerjee", "Singh", "Khan", "Fernandes", "Rao", "Joshi", "Chatterjee", "Varghese", "Bhat"]


def person(used):
    while True:
        name = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
        if name not in used:
            used.add(name)
            return name


def phone():
    return f"+91 9{rng.randint(100000000, 999999999)}"


def messy_key(key):
    return rng.choice([key, key, key, f" {key}", key.lower(), f"{key} "])


def messy_date(d):
    kind = rng.randrange(4)
    if kind == 0:
        return dt.datetime(d.year, d.month, d.day)  # real Excel date cell
    if kind == 1:
        return d.isoformat()
    if kind == 2:
        return d.strftime("%d/%m/%Y")
    return d.strftime("%b %d, %Y")


def messy_money(x):
    return rng.choice([f"₹{x:,.2f}", f"INR {x:.0f}", x, f"{x:,.2f}"])


def messy_int(n):
    return rng.choice([n, f"{n:,}", str(n)])


def messy_bool(b):
    return rng.choice(["Y", "yes", "TRUE", "1"] if b else ["N", "no", "FALSE", "0"])


def rand_date(start, end):
    return start + dt.timedelta(days=rng.randrange((end - start).days + 1))


# --------------------------------------------------------------------------- retail model
def build_retail():
    people = set()
    planted = {  # name, rating
        "SUP-001": ("Bluepeak Textiles", 4.1),
        "SUP-002": ("Kaveri Packaging", 3.6),
        "SUP-003": ("Orion Components", 3.9),
        "SUP-004": ("Suvarna Foods", 4.4),
    }
    foreign = [("Müller Präzisionsteile GmbH", "Germany"), ("Société Lumière SARL", "France"),
               ("Østergaard Tools ApS", "Denmark")]
    prefixes = ["Apex", "Nimbus", "Sahyadri", "Coromandel", "Deccan", "Indus", "Kestrel", "Lotus", "Meridian",
                "Narmada", "Pinnacle", "Quartz", "Riverstone", "Saffron", "Terra", "Vega", "Zenith", "Ganga",
                "Himalaya", "Konkan"]
    suffixes = ["Traders", "Industries", "Polymers", "Logistics", "Electricals", "Hardware", "Components",
                "Foods", "Textiles", "Packaging"]
    names = set(n for n, _ in planted.values()) | {n for n, _ in foreign}
    generated = []
    while len(generated) < 41:
        n = f"{rng.choice(prefixes)} {rng.choice(suffixes)}"
        if n not in names:
            names.add(n)
            generated.append((n, "India"))
    suppliers = {}
    others = foreign + generated
    rng.shuffle(others)
    for i in range(1, 49):
        sid = f"SUP-{i:03d}"
        if sid in planted:
            name, rating, country = *planted[sid], "India"
        else:
            name, country = others.pop()
            rating = None if i in (17, 29, 40) else round(rng.uniform(3.0, 5.0), 1)
        contact = person(people)
        suppliers[sid] = {
            "supplier_id": sid, "name": name, "country": country, "rating": rating,
            "contact_person": contact,
            "contact_email": f"{contact.lower().replace(' ', '.')}@{name.split()[0].lower()}.example",
            "contact_phone": phone(),
            "bank_account": f"{rng.randint(10**11, 10**12 - 1)} / IFSC {rng.choice(['HDFC', 'ICIC', 'SBIN', 'UTIB'])}0{rng.randint(100000, 999999)}",
            "onboarded": rand_date(dt.date(2019, 1, 1), dt.date(2026, 6, 30)),
        }

    warehouses = {}
    for code, city, cap in [("WH-CHN", "Chennai", 4000), ("WH-MUM", "Mumbai", 25000), ("WH-DEL", "Delhi", 24500),
                            ("WH-BLR", "Bangalore", 18000), ("WH-HYD", "Hyderabad", 15000),
                            ("WH-KOL", "Kolkata", 12000), ("WH-PUN", "Pune", 14000),
                            ("WH-AMD", "Ahmedabad", 11000), ("WH-KOC", "Kochi", 6500)]:
        warehouses[code] = {"warehouse_id": code, "city": city, "capacity": cap,
                            "manager": person(people), "manager_phone": phone()}

    # Products: the four planted suppliers get fixed counts, the rest spread over the other 44.
    categories = ["Apparel", "Home", "Kitchen", "Electronics", "Grocery", "Packaging", "Tools", "Stationery"]
    nouns = ["Cotton Towel", "Steel Tumbler", "LED Bulb", "Basmati Rice", "Corrugated Box", "Hex Key Set",
             "Notebook", "Bedsheet", "Pressure Cooker", "USB Cable", "Masala Mix", "Bubble Wrap", "Drill Bit",
             "Gel Pen", "Curtain", "Tawa", "Extension Board", "Tea Dust", "Packing Tape", "Spanner"]
    owners = (["SUP-001"] * 45 + ["SUP-002"] * 14 + ["SUP-003"] * 11 + ["SUP-004"] * 18)
    rest = [f"SUP-{i:03d}" for i in range(5, 49)]
    owners += rest + [rng.choice(rest) for _ in range(312 - len(owners) - len(rest))]
    products = {}
    for n, owner in enumerate(owners):
        sku = f"SKU-{10001 + n}"
        cat = rng.choice(categories)
        if rng.random() < 0.08:
            cat = f"{cat}; {rng.choice(categories)}"
        products[sku] = {"sku": sku, "name": f"{rng.choice(nouns)} {rng.choice(['Classic', 'Pro', 'Lite', 'XL', 'Eco'])}",
                         "category": cat, "unit_price": round(rng.uniform(20, 4000), 2),
                         "supplier_id": owner, "discontinued": rng.random() < 0.1, "substitute": None}
    skus = list(products)
    # 5 unrecoverable product rows: supplier written as a name, not an ID (not from planted suppliers)
    bad_products = rng.sample([s for s in skus if products[s]["supplier_id"] not in planted], 5)
    loadable = [s for s in skus if s not in bad_products]
    # substitutes, including forward references (SKU-10007 -> SKU-10213)
    products["SKU-10007"]["substitute"] = "SKU-10213"
    for sku in rng.sample(loadable, 30):
        if sku != "SKU-10007":
            products[sku]["substitute"] = rng.choice([s for s in loadable if s != sku])

    # Inventory: Chennai stocks only 37/12/9/15 SKUs of the four planted suppliers.
    inventory = []  # (sku, wh, qty, date) in file order; last row wins
    by_sup = defaultdict(list)
    for s in loadable:
        by_sup[products[s]["supplier_id"]].append(s)
    for sid, k in [("SUP-001", 37), ("SUP-002", 12), ("SUP-003", 9), ("SUP-004", 15)]:
        for sku in by_sup[sid][:k]:
            inventory.append([sku, "WH-CHN", rng.randint(5, 120), rand_date(dt.date(2026, 8, 1), dt.date(2026, 9, 20))])
    chennai = {r[0] for r in inventory}
    non_chennai_pool = [s for s in loadable]
    for wh in warehouses:
        if wh == "WH-CHN":
            continue
        for sku in rng.sample(non_chennai_pool, 120):
            inventory.append([sku, wh, rng.randint(10, 900), rand_date(dt.date(2026, 8, 1), dt.date(2026, 9, 20))])
    for wh in ("WH-MUM", "WH-PUN"):
        if not any(r[0] == "SKU-10001" and r[1] == wh for r in inventory):
            inventory.append(["SKU-10001", wh, rng.randint(50, 400), dt.date(2026, 9, 1)])
    # 10 repeated (sku, wh) rows with a new count, placed later in the sheet
    for r in rng.sample(inventory, 9) + [next(r for r in inventory if r[0] == "SKU-10001" and r[1] == "WH-MUM")]:
        inventory.append([r[0], r[1], r[2] + rng.randint(1, 60), dt.date(2026, 9, 25)])
    rng.shuffle(inventory)
    # keep the SKU-10001/WH-MUM correction as the last occurrence of that pair
    fix = [r for r in inventory if r[0] == "SKU-10001" and r[1] == "WH-MUM"]
    for r in fix:
        inventory.remove(r)
    inventory.extend(sorted(fix, key=lambda r: r[3]))

    customers = {}
    regions = ["South", "North", "West", "East"]
    for i in range(1, 121):
        cid = f"CUS-{i:04d}"
        name = person(people)
        customers[cid] = {
            "customer_id": cid, "name": name,
            "email": f"{name.lower().replace(' ', '.')}{rng.randint(1, 99)}@mail.example",
            "phone": phone(),
            "dob": None if rng.random() < 0.1 else rand_date(dt.date(1960, 1, 1), dt.date(2004, 12, 31)),
            "pan": "".join(rng.choice("ABCDEFGHJKLMNPQRSTUVWXYZ") for _ in range(5)) + f"{rng.randint(1000, 9999)}" + rng.choice("ABCDEFGHJKLMNPQRSTUVWXYZ"),
            "address": f"{rng.randint(1, 250)}, {rng.choice(['MG Road', 'Anna Salai', 'Park Street', 'Linking Road', 'FC Road'])}, {rng.choice([w['city'] for w in warehouses.values()])}",
            "region": rng.choice(regions),
        }

    # Orders: one row per order line.
    status_w = [("Delivered", 60), ("Shipped", 20), ("Pending", 12), ("Cancelled", 8)]
    cust_ids = list(customers)
    cust_w = [12 if c == "CUS-0042" else 1 for c in cust_ids]
    orders, lines = {}, []
    order_skus = [s for s in loadable]
    n = 0
    while len(lines) < 4812 - 27:
        n += 1
        oid = f"ORD-2026-{n:05d}"
        cid = rng.choices(cust_ids, cust_w)[0]
        odate = rand_date(dt.date(2026, 6, 1), dt.date(2026, 9, 20))
        status = rng.choices([s for s, _ in status_w], [w for _, w in status_w])[0]
        shipped = odate + dt.timedelta(days=rng.randint(1, 5)) if status in ("Delivered", "Shipped") else None
        orders[oid] = {"order_id": oid, "date": odate, "status": status, "customer_id": cid,
                       "warehouse": rng.choice(list(warehouses)), "shipped_on": shipped}
        for sku in rng.sample(order_skus, rng.randint(1, 5)):
            if len(lines) >= 4812 - 27:
                break
            qty = rng.randint(10, 40) if cid == "CUS-0042" else rng.randint(1, 20)
            lines.append({"order_id": oid, "sku": sku, "qty": qty})

    return {"suppliers": suppliers, "warehouses": warehouses, "products": products,
            "bad_products": bad_products, "loadable_products": loadable, "inventory": inventory,
            "chennai": chennai, "customers": customers, "orders": orders, "lines": lines}


def note_text():
    return rng.choice([
        "Leave at security gate", "Fragile - handle with care", "Partial delivery accepted",
        f"Customer asked to call {rng.choice(FIRST)} on {phone()}",  # PII hiding in free text
        f"Deliver to {rng.randint(1, 99)}, Anna Salai, flat {rng.randint(1, 30)}",
        "Invoice to be emailed", "Reschedule if raining",
    ])


def write_retail_xlsx(m, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    ws.append(["Supplier Master — exported 22/09/2026 by procurement"])
    ws.append([])
    ws.append(["Supplier ID", "Supplier Name", "Country", "Rating (1-5)", "Contact Person", "Contact Email",
               "Contact Phone", "Bank Account", "Onboarded"])
    rows = list(m["suppliers"].values())
    rows += [dict(m["suppliers"]["SUP-011"]), dict(m["suppliers"]["SUP-023"])]  # duplicates
    for s in rows:
        rating = "N/A" if s["rating"] is None else rng.choice([s["rating"], str(s["rating"])])
        ws.append([messy_key(s["supplier_id"]), s["name"], s["country"], rating, s["contact_person"],
                   s["contact_email"], s["contact_phone"], s["bank_account"], messy_date(s["onboarded"])])

    ws = wb.create_sheet("Products")
    ws.append(["SKU", "Product", "Category", "Unit Price", "Supplier", "Substitute SKU", "Discontinued"])
    for sku, p in m["products"].items():
        sup = m["suppliers"][p["supplier_id"]]["name"] if sku in m["bad_products"] else messy_key(p["supplier_id"])
        ws.append([messy_key(sku), p["name"], p["category"], messy_money(p["unit_price"]), sup,
                   p["substitute"] or "", messy_bool(p["discontinued"])])

    ws = wb.create_sheet("Warehouses")
    ws.append(["WH Code", "City", "Capacity (units)", "Manager", "Manager Phone"])
    for w in m["warehouses"].values():
        ws.append([w["warehouse_id"], w["city"], messy_int(w["capacity"]), w["manager"], w["manager_phone"]])

    ws = wb.create_sheet("Customers")
    ws.append(["customer_id", "Name", "Email", "Phone", "DOB", "PAN", "Address", "Region"])
    rows = list(m["customers"].values()) + [m["customers"][c] for c in ("CUS-0007", "CUS-0063", "CUS-0101")]
    for c in rows:
        ws.append([messy_key(c["customer_id"]), c["name"], c["email"], c["phone"],
                   messy_date(c["dob"]) if c["dob"] else "", c["pan"], c["address"], c["region"]])

    ws = wb.create_sheet("Orders")
    ws.append(["order_id", "order_date", "customer_id", "sku", "qty", "warehouse", "shipped_on", "status",
               "notes", "row_no"])
    out = []
    for ln in m["lines"]:
        o = m["orders"][ln["order_id"]]
        out.append([ln["order_id"], messy_date(o["date"]), messy_key(o["customer_id"]), messy_key(ln["sku"]),
                    ln["qty"], o["warehouse"], messy_date(o["shipped_on"]) if o["shipped_on"] else "", o["status"],
                    note_text() if rng.random() < 0.02 else ""])
    bad = []
    reasons = ([("missing_customer_id", 9), ("unknown_sku", 7), ("unknown_warehouse", 5), ("missing_order_id", 6)])
    k = 90000
    for reason, count in reasons:
        for _ in range(count):
            k += 1
            row = [f"ORD-2026-{k}", messy_date(dt.date(2026, 9, 1)), rng.choice(list(m["customers"])),
                   rng.choice(m["loadable_products"]), rng.randint(1, 9), "WH-MUM", "", "Pending", ""]
            if reason == "missing_customer_id":
                row[2] = ""
            elif reason == "unknown_sku":
                row[3] = rng.choice(["SKU-00000", "SKU-77777", "SKU-1O001"])  # letter O, not zero
            elif reason == "unknown_warehouse":
                row[5] = rng.choice(["WH-99", "WH-NGP"])
            else:
                row[0] = ""
            bad.append(row)
    for row in bad:
        out.insert(rng.randrange(len(out)), row)
    for i, row in enumerate(out, 1):
        ws.append(row + [i])
        if i % 800 == 0:
            ws.append([])  # stray blank rows
    ws = wb.create_sheet("Inventory")
    ws.append(["sku", "wh_code", "stock_qty", "last_counted"])
    for sku, wh, qty, d in m["inventory"]:
        ws.append([messy_key(sku), wh, qty, messy_date(d)])
    wb.save(path)
    return Counter(r for r, c in reasons for _ in range(c))


# --------------------------------------------------------------------------- add-data files
def write_october(m, path):
    header = ["SKU", "Order ID", "Qty", "Customer ID", "Order Date", "Status", "Warehouse", "Shipped On", "Notes"]
    new_orders, new_lines = {}, []
    n = 10000
    while len(new_lines) < 1200:
        n += 1
        oid = f"ORD-2026-{n:05d}"
        odate = rand_date(dt.date(2026, 10, 1), dt.date(2026, 10, 28))
        status = rng.choices(["Delivered", "Shipped", "Pending", "Cancelled"], [55, 25, 12, 8])[0]
        new_orders[oid] = {"order_id": oid, "date": odate, "status": status,
                           "customer_id": rng.choice(list(m["customers"])), "warehouse": rng.choice(list(m["warehouses"])),
                           "shipped_on": odate + dt.timedelta(days=2) if status in ("Delivered", "Shipped") else None}
        for sku in rng.sample(m["loadable_products"], rng.randint(1, 4)):
            if len(new_lines) < 1200:
                new_lines.append({"order_id": oid, "sku": sku, "qty": rng.randint(1, 20)})

    def row(o, ln):
        return {"SKU": ln["sku"], "Order ID": o["order_id"], "Qty": ln["qty"], "Customer ID": o["customer_id"],
                "Order Date": messy_date(o["date"]), "Status": o["status"], "Warehouse": o["warehouse"],
                "Shipped On": messy_date(o["shipped_on"]) if o["shipped_on"] else "", "Notes": ""}

    rows = [row(new_orders[ln["order_id"]], ln) for ln in new_lines]
    rows += [row(m["orders"][ln["order_id"]], ln) for ln in rng.sample(m["lines"], 20)]  # already loaded
    bad = Counter()
    for reason, count in [("unknown_customer_id", 4), ("unknown_sku", 3), ("missing_order_id", 3)]:
        for i in range(count):
            r = row({"order_id": f"ORD-2026-{19900 + len(rows)}", "date": dt.date(2026, 10, 15), "status": "Pending",
                     "customer_id": "CUS-0005", "warehouse": "WH-DEL", "shipped_on": None},
                    {"sku": "SKU-10020", "qty": 3})
            if reason == "unknown_customer_id":
                r["Customer ID"] = f"CUS-09{90 + i}"
            elif reason == "unknown_sku":
                r["SKU"] = f"SKU-2{i}000"
            else:
                r["Order ID"] = ""
            rows.append(r)
            bad[reason] += 1
    rng.shuffle(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Oct Orders"
    ws.append(header)
    for r in rows:
        ws.append([r[h] for h in header])
    wb.save(path)
    return new_orders, new_lines, bad


def write_inventory_update(m, path):
    existing = {}
    for sku, wh, qty, _ in m["inventory"]:
        existing[(sku, wh)] = qty
    updates = rng.sample(sorted(k for k in existing if k != ("SKU-10001", "WH-CHN")), 29) + [("SKU-10001", "WH-CHN")]
    new_pairs = []
    while len(new_pairs) < 14:
        k = (rng.choice(m["loadable_products"]), rng.choice([w for w in m["warehouses"] if w != "WH-CHN"]))
        if k not in existing and k not in new_pairs:
            new_pairs.append(k)
    rows = [[sku, wh, rng.randint(0, 500), "2026-09-28"] for sku, wh in updates + new_pairs]
    rows += [["", rng.choice(list(m["warehouses"])), rng.randint(1, 50), "2026-09-28"] for _ in range(12)]
    rng.shuffle(rows)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM + ; + CRLF
        w = csv.writer(f, delimiter=";", lineterminator="\r\n")
        w.writerow(["sku", "wh_code", "stock_qty", "last_counted"])
        w.writerows(rows)
    return rows


# --------------------------------------------------------------------------- finance ledger (second KB)
def write_finance(path):
    people = set()
    approvers = [person(people) for _ in range(6)]
    vendors = ["Bluepeak Textiles", "Tata Power", "Airtel Business", "Infosys BPM", "Deccan Couriers",
               "Sodexo Facilities", "Amazon Web Services", "Kaveri Packaging"]
    accounts = [("5100", "Utilities"), ("5200", "Telecom"), ("5300", "Cloud Hosting"), ("5400", "Outsourcing"),
                ("5500", "Logistics"), ("5600", "Facilities"), ("5700", "Raw Material"), ("2100", "Accounts Payable")]
    ccs = ["CC-OPS", "CC-MKT", "CC-IT", "CC-HR", "CC-FIN"]
    rows, truth = [], []
    for i in range(1, 481):
        acct = rng.choice(accounts)
        amt = round(rng.uniform(1500, 250000), 2)
        if rng.random() < 0.12:
            amt = -amt  # credit notes
        vendor = "Amazon Web Services" if acct[0] == "5300" else rng.choice(vendors)
        if acct[0] == "5300":
            amt = abs(amt) + 150000  # AWS is the biggest spender by construction
        appr = rng.choice(approvers)
        e = {"entry": f"JE-{2026}{i:05d}", "date": rand_date(dt.date(2026, 4, 1), dt.date(2026, 9, 25)),
             "account": acct[0], "account_name": acct[1], "cc": rng.choice(ccs), "vendor": vendor,
             "amount": amt, "approver": appr}
        truth.append(e)
        amount_txt = f"({abs(amt):,.2f})" if amt < 0 else f"{amt:,.2f}"
        rows.append([e["entry"], e["date"].strftime("%d/%m/%Y"), e["account"], e["account_name"], e["cc"], vendor,
                     amount_txt, "INR", appr, f"{appr.lower().replace(' ', '.')}@graphbase-retail.example"])
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\r\n")
        w.writerow(["Entry No", "Posting Date", "GL Account", "Account Name", "Cost Centre", "Vendor", "Amount",
                    "Currency", "Approved By", "Approver Email"])
        w.writerows(rows)
    return truth


# --------------------------------------------------------------------------- RAG documents
FILLER = [
    "All staff must follow this policy consistently across stores, warehouses and the online channel. Exceptions "
    "require written approval from the regional operations head and must be logged in the exceptions register.",
    "Store teams should explain the policy politely to customers and offer an exchange before a refund where the "
    "product is in resaleable condition. Refunds are issued to the original payment method.",
    "This document is reviewed every year by the Customer Experience and Legal teams. Printed copies may be out of "
    "date; the version on the intranet is authoritative.",
]


def write_returns_pdf(path):
    styles = getSampleStyleSheet()
    story = [Paragraph("Customer Returns Policy", styles["Title"]),
             Paragraph("Version 4.2, effective 1 January 2025", styles["Normal"]), Spacer(1, 12),
             Paragraph("1. Standard return window", styles["Heading2"]),
             Paragraph("Unless a category rule below says otherwise, customers may return products within "
                       "<b>30 days</b> of delivery with proof of purchase.", styles["Normal"])]
    story += [Paragraph(t, styles["Normal"]) for t in FILLER]
    story += [Paragraph("2. Category rules", styles["Heading2"]),
              Paragraph("The table below overrides the standard window for specific categories.", styles["Normal"]),
              Spacer(1, 6)]
    table = Table([["Category", "Return window", "Restocking fee"],
                   ["Electronics", "15 days", "12% of item price"],
                   ["Apparel", "30 days", "None"],
                   ["Furniture", "30 days", "8% of item price"],
                   ["Perishables / Grocery", "Not returnable", "-"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story += [table, Spacer(1, 12)]
    story += [Paragraph(t, styles["Normal"]) for t in FILLER] + [PageBreak()]
    story += [Paragraph("3. Refund timelines", styles["Heading2"]),
              Paragraph("Approved refunds are processed within 7 working days. UPI and card refunds may take a "
                        "further 3 to 5 bank working days to appear.", styles["Normal"]),
              Paragraph("4. Contact", styles["Heading2"]),
              Paragraph("Returns desk: returns@graphbase-retail.example, phone +91 44 4000 1234 "
                        "(Mon to Sat, 9:00 to 18:00). Escalations go to Deepa Menon, Customer Care Manager.",
                        styles["Normal"]),
              Paragraph("5. History", styles["Heading2"]),
              Paragraph("Before 1 January 2025 the standard return window was 45 days. That rule is superseded "
                        "and must not be quoted to customers.", styles["Normal"])]
    story += [Paragraph(t, styles["Normal"]) for t in FILLER]

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(40, 20, f"Graphbase Retail - Internal - Page {doc.page}")
        canvas.drawString(40, A4[1] - 25, "CONFIDENTIAL: Customer Returns Policy v4.2")

    SimpleDocTemplate(str(path), pagesize=A4).build(story, onFirstPage=footer, onLaterPages=footer)


def write_handbook_docx(path):
    d = Document()
    d.add_heading("Vendor Handbook", 0)
    d.add_paragraph("This handbook applies to every supplier registered with Graphbase Retail.")
    d.add_heading("Onboarding", 1)
    for item in ["GST registration certificate", "Cancelled cheque for the payout account",
                 "Signed code of conduct", "Two trade references"]:
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("Payment terms", 1)
    d.add_paragraph("Invoices are paid on net 45 day terms from the date a correct invoice is received. "
                    "Early payment at a 1.5% discount is available on request for invoices paid within 10 days.")
    d.add_heading("Performance targets", 1)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "KPI", "Target"
    for k, v in [("On time in full (OTIF)", "95% or higher"), ("Defect rate", "below 1.5%"),
                 ("Invoice accuracy", "98% or higher")]:
        c = t.add_row().cells
        c[0].text, c[1].text = k, v
    d.add_heading("Late delivery penalty", 1)
    d.add_paragraph("Late deliveries attract a penalty of 2% of the purchase order value for each week of delay, "
                    "capped at 10% of the purchase order value.")
    d.add_heading("Escalation contacts", 1)
    d.add_paragraph("Procurement lead: Priya Nair, priya.nair@graphbase-retail.example, +91 98400 11223.")
    d.add_paragraph("Accounts payable: Venkat Rao, ap@graphbase-retail.example, +91 98400 44556.")
    for t_ in FILLER:
        d.add_paragraph(t_)
    d.save(path)


def write_sop_txt(path):
    text = """WAREHOUSE STANDARD OPERATING PROCEDURES (Chennai, WH-CHN)
=========================================================

1. Cold storage
The Chennai warehouse has a cold room that must be held between 2 and 8 degrees Celsius.
குளிர் அறை வெப்பநிலை 2 முதல் 8 டிகிரி செல்சியஸ் வரை இருக்க வேண்டும்.
Temperature is logged every 4 hours. Two consecutive readings outside the range must be escalated to the shift lead.

2. Dock hours
Inbound docks are open 06:00 to 14:00, Monday to Saturday. Outbound dispatch runs 14:00 to 22:00.
Vehicles arriving outside dock hours wait in the holding bay.

3. Equipment
Forklift operators must renew their certification every 12 months.
Fire drills are held once a quarter; the last drill was on 12 August 2026.

4. Contacts
Shift lead (day): Karthik Pillai, +91 94440 55667
"""
    path.write_bytes(text.replace("\n", "\r\n").encode("utf-8"))


# --------------------------------------------------------------------------- main
def main():
    SAMPLES.mkdir(exist_ok=True)
    EXPECTED.mkdir(exist_ok=True)
    (SAMPLES / "edge_cases").mkdir(exist_ok=True)
    m = build_retail()
    bad_orders = write_retail_xlsx(m, SAMPLES / "supplier_orders.xlsx")
    oct_orders, oct_lines, oct_bad = write_october(m, SAMPLES / "supplier_orders_october.xlsx")
    inv_update = write_inventory_update(m, SAMPLES / "warehouse_update.csv")
    ledger = write_finance(SAMPLES / "finance_ledger.csv")
    write_returns_pdf(SAMPLES / "returns_policy.pdf")
    write_handbook_docx(SAMPLES / "vendor_handbook.docx")
    write_sop_txt(SAMPLES / "warehouse_sop.txt")
    (SAMPLES / "edge_cases" / "corrupt.xlsx").write_bytes(bytes(rng.randrange(256) for _ in range(4096)))
    (SAMPLES / "edge_cases" / "empty.csv").write_text("sku,wh_code,stock_qty\n")

    # ---- ground truth after the initial build
    stock = {}
    for sku, wh, qty, _ in m["inventory"]:
        stock[(sku, wh)] = qty
    qty_by_cust = Counter()
    for ln in m["lines"]:
        qty_by_cust[m["orders"][ln["order_id"]]["customer_id"]] += ln["qty"]
    (top_c, top_q), (_, second_q) = qty_by_cust.most_common(2)
    assert top_q > second_q
    chennai_sups = sorted({m["suppliers"][m["products"][s]["supplier_id"]]["name"] for s in m["chennai"]})
    low_rated = sorted(m["suppliers"][sid]["name"] for sid in ("SUP-002", "SUP-003"))
    cancelled = sum(o["status"] == "Cancelled" for o in m["orders"].values())
    sku1_stock = sum(q for (s, _), q in stock.items() if s == "SKU-10001")
    blr = sum(1 for (_, wh) in stock if wh == "WH-BLR")
    subs = sum(1 for s in m["loadable_products"] if m["products"][s]["substitute"])

    # ---- after add-data
    stock_after = dict(stock)
    for sku, wh, qty, _ in inv_update:
        if sku:
            stock_after[(sku, wh)] = qty
    all_orders = {**m["orders"], **oct_orders}
    cc_it = round(sum(e["amount"] for e in ledger if e["cc"] == "CC-IT"), 2)
    spend = Counter()
    for e in ledger:
        spend[e["vendor"]] += e["amount"]

    manifest = {
        "seed": 20260929,
        "ingest_contract": __doc__.split("contract the backend must honour")[1].split('"""')[0].strip(),
        "files": {
            "supplier_orders.xlsx": {
                "kb": "retail_supply_chain_kg", "kind": "graph_initial",
                "sheets": {"Suppliers": {"header_row": 3, "data_rows": 50},
                           "Products": {"header_row": 1, "data_rows": 312},
                           "Warehouses": {"header_row": 1, "data_rows": 9},
                           "Customers": {"header_row": 1, "data_rows": 123},
                           "Orders": {"header_row": 1, "data_rows": 4812, "blank_rows": 6},
                           "Inventory": {"header_row": 1, "data_rows": len(m["inventory"])}},
                "expected_nodes": {  # keyed by source sheet + key column (labels are chosen by the LLM)
                    "Suppliers.Supplier ID": 48, "Products.SKU": 307, "Warehouses.WH Code": 9,
                    "Customers.customer_id": 120, "Orders.order_id": len(m["orders"])},
                "expected_relationships": {
                    "supplier->product": 307, "product->substitute_product": subs,
                    "product->warehouse (stock)": len(stock), "customer->order": len(m["orders"]),
                    "order->product (line)": len(m["lines"]), "order->warehouse": len(m["orders"])},
                "expected_rejections": {"Products": {"supplier_given_as_name": 5},
                                        "Orders": dict(bad_orders)},
                "pii": {
                    "Suppliers": {"Contact Person": "person_name", "Contact Email": "email",
                                  "Contact Phone": "phone", "Bank Account": "bank_account"},
                    "Warehouses": {"Manager": "person_name", "Manager Phone": "phone"},
                    "Customers": {"Name": "person_name", "Email": "email", "Phone": "phone",
                                  "DOB": "date_of_birth", "PAN": "government_id", "Address": "address"},
                    "Orders": {"notes": "free_text_may_contain_pii"}},
                "not_pii_traps": ["Suppliers.Supplier Name (company)", "Products.Product", "Warehouses.City",
                                  "Customers.Region", "Orders.row_no"],
                "junk_columns": ["Orders.row_no"],
            },
            "supplier_orders_october.xlsx": {
                "kb": "retail_supply_chain_kg", "kind": "graph_add_data", "rows": 1230,
                "header_note": "same columns as Orders but renamed/reordered, no row_no",
                "expected": {"new_orders": len(oct_orders), "new_order_lines": len(oct_lines),
                             "already_loaded_lines": 20, "rejected": dict(oct_bad)}},
            "warehouse_update.csv": {
                "kb": "retail_supply_chain_kg", "kind": "graph_add_data", "rows": 56,
                "format": "UTF-8 BOM, ';' delimiter, CRLF",
                "expected": {"updated_pairs": 30, "new_pairs": 14, "rejected": {"missing_sku": 12}}},
            "finance_ledger.csv": {
                "kb": "finance_ledger_kg", "kind": "graph_initial", "rows": 480,
                "format": "UTF-8 BOM, ';' delimiter, CRLF, accounting negatives '(1,200.00)', DD/MM/YYYY",
                "pii": {"Approved By": "person_name", "Approver Email": "email"},
                "isolation_trap": "vendors 'Bluepeak Textiles' and 'Kaveri Packaging' also exist in the retail KB"},
            "returns_policy.pdf": {"kb": "retail_policies_rag", "kind": "rag",
                                   "pii": {"email": 1, "phone": 1, "person_name": 1},
                                   "not_pii_traps": ["returns@ is a shared mailbox but still counts as an email"]},
            "vendor_handbook.docx": {"kb": "retail_policies_rag", "kind": "rag",
                                     "pii": {"person_name": 2, "email": 2, "phone": 2}},
            "warehouse_sop.txt": {"kb": "retail_policies_rag", "kind": "rag", "format": "CRLF, contains Tamil",
                                  "pii": {"person_name": 1, "phone": 1}},
            "edge_cases/corrupt.xlsx": {"expect": "clear error, no job crash"},
            "edge_cases/empty.csv": {"expect": "clear error: no data rows"},
        },
        "questions": [
            {"kb": "retail_supply_chain_kg", "level": "core", "q": "How many suppliers are there?", "answer": 48},
            {"kb": "retail_supply_chain_kg", "level": "core",
             "q": "Which suppliers deliver products stored in the Chennai warehouse?", "answer": chennai_sups},
            {"kb": "retail_supply_chain_kg", "level": "core", "follow_up": True,
             "q": "Which of those have a rating below 4?", "answer": low_rated},
            {"kb": "retail_supply_chain_kg", "level": "core",
             "q": "Which customer ordered the largest total quantity?",
             "answer": m["customers"][top_c]["name"], "detail": {"customer_id": top_c, "qty": top_q}},
            {"kb": "retail_supply_chain_kg", "level": "core", "q": "How many orders were cancelled?",
             "answer": cancelled},
            {"kb": "retail_supply_chain_kg", "level": "core", "q": "Which warehouse has the largest capacity?",
             "answer": "Mumbai"},
            {"kb": "retail_supply_chain_kg", "level": "core", "q": "What is the substitute product for SKU-10007?",
             "answer": "SKU-10213"},
            {"kb": "retail_supply_chain_kg", "level": "core",
             "q": "What is the total stock of SKU-10001 across all warehouses?", "answer": sku1_stock,
             "after_add_data": sum(q for (s, _), q in stock_after.items() if s == "SKU-10001")},
            {"kb": "retail_supply_chain_kg", "level": "core", "q": "How many products does Bluepeak Textiles supply?",
             "answer": 45},
            {"kb": "retail_supply_chain_kg", "level": "stretch",
             "q": "How many products are stored in the Bengaluru warehouse?", "answer": blr,
             "trap": "city is stored as 'Bangalore'"},
            {"kb": "retail_supply_chain_kg", "level": "core", "after_add_data_only": True,
             "q": "How many orders were cancelled?",
             "answer": sum(o["status"] == "Cancelled" for o in all_orders.values())},
            {"kb": "finance_ledger_kg", "level": "core", "q": "What is the total amount posted to cost centre CC-IT?",
             "answer": cc_it, "tolerance": 0.01},
            {"kb": "finance_ledger_kg", "level": "core", "q": "Which vendor has the highest total spend?",
             "answer": spend.most_common(1)[0][0]},
            {"kb": "retail_policies_rag", "level": "core", "q": "What is the standard return window?",
             "answer": "30 days", "trap": "superseded 45-day rule is also in the document"},
            {"kb": "retail_policies_rag", "level": "core", "q": "What is the restocking fee for electronics?",
             "answer": "12%", "trap": "fact is only in a PDF table"},
            {"kb": "retail_policies_rag", "level": "core", "q": "What are the vendor payment terms?",
             "answer": "net 45 days"},
            {"kb": "retail_policies_rag", "level": "core", "q": "What is the late delivery penalty?",
             "answer": "2% of PO value per week of delay, capped at 10%"},
            {"kb": "retail_policies_rag", "level": "core", "q": "What temperature must the Chennai cold room be kept at?",
             "answer": "2 to 8 degrees Celsius"},
            {"kb": "retail_policies_rag", "level": "core", "q": "What are the inbound dock hours?",
             "answer": "06:00 to 14:00, Monday to Saturday"},
            {"kb": "retail_supply_chain_kg", "level": "security", "as_user": "unauthorised",
             "q": "How many suppliers are there?", "answer": "403 access denied"},
            {"kb": "retail_supply_chain_kg", "level": "security",
             "q": "Delete all suppliers", "answer": "refused: chat Cypher is read-only"},
            {"kb": "finance_ledger_kg", "level": "security",
             "q": "Which products are stored in the Chennai warehouse?",
             "answer": "no retail data (KB isolation)"},
        ],
    }
    (EXPECTED / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False, default=str))
    print(json.dumps({k: v.get("expected_nodes", v.get("expected", "")) for k, v in manifest["files"].items()
                      if isinstance(v, dict) and ("expected_nodes" in v or "expected" in v)}, indent=1, default=str))


if __name__ == "__main__":
    main()
