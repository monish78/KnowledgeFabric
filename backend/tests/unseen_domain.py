"""A workbook from a domain the code has never seen (a school), for generality tests."""
import random

from openpyxl import Workbook


def make_school_workbook(path):
    rng = random.Random(7)
    first = ["Asha", "Rahul", "Neha", "Imran", "Lakshmi", "Tom", "Fatima", "Wei", "Carlos", "Anita", "Dev", "Sara",
             "Vikram", "Meena", "John", "Priti"]
    last = ["Kumar", "Shah", "Das", "Brown", "Iyer", "Khan", "Lopez", "Chen", "Nair", "Joshi", "Roy", "Patel"]
    names = [f"{f} {l}" for f in first for l in last]
    rng.shuffle(names)
    wb = Workbook()
    ws = wb.active
    ws.title = "Teachers"
    ws.append(["Teacher Code", "Full Name", "Subject Area", "Joined"])
    teachers = [f"T{i:03d}" for i in range(1, 13)]
    for t in teachers:
        ws.append([t, names.pop(), rng.choice(["Maths", "Science", "Arts"]), f"2020-0{rng.randint(1, 9)}-15"])
    ws = wb.create_sheet("Courses")
    ws.append(["course_code", "Title", "Credits", "Taught By"])
    courses = [f"C-{i:02d}" for i in range(1, 21)]
    for c in courses:
        ws.append([c, f"Course {c}", rng.choice([2, 3, 4]), rng.choice(teachers)])
    ws = wb.create_sheet("Students")
    ws.append(["Roll No", "Student Name", "Email", "Year"])
    students = [f"S{i:04d}" for i in range(1, 151)]
    for s in students:
        n = names.pop()
        ws.append([s, n, f"{n.lower().replace(' ', '.')}@school.example", rng.choice([1, 2, 3, 4])])
    ws = wb.create_sheet("Enrolments")
    ws.append(["roll_no", "course_code", "Grade", "Term"])
    pairs = set()
    while len(pairs) < 600:
        pairs.add((rng.choice(students), rng.choice(courses)))
    pairs = sorted(pairs)
    for s, c in pairs:
        ws.append([s, c, rng.choice(["A", "B", "C", "D"]), rng.choice(["2025-S1", "2025-S2"])])
    ws.append(["S9999", "C-01", "A", "2025-S1"])  # unknown student -> rejected
    wb.save(path)
    return {"teachers": 12, "courses": 20, "students": 150, "enrolments": 600, "taught_by": 20, "rejected": 1}
